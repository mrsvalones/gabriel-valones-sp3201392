from io import BytesIO
from datetime import datetime, time, timedelta, timezone
from zoneinfo import ZoneInfo
from decimal import Decimal
from flask import Blueprint,g,request,render_template,send_file
from sqlalchemy import select,func
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from ..models import *
from ..security import login_required,admin_required,reauthenticate
from ..services.operations import number,audit

bp=Blueprint('reports',__name__)
TZ=ZoneInfo('America/Sao_Paulo')

def period():
    today=datetime.now(TZ).date()
    try:
        start=datetime.strptime(request.args.get('start',str(today.replace(day=1))),'%Y-%m-%d').date()
        end=datetime.strptime(request.args.get('end',str(today)),'%Y-%m-%d').date()
    except ValueError: raise ValueError('Datas inválidas.')
    if end<start or (end-start).days>366: raise ValueError('Período deve ter de 1 a 367 dias.')
    a=datetime.combine(start,time.min,TZ); b=datetime.combine(end+timedelta(days=1),time.min,TZ)
    return start,end,a.astimezone(timezone.utc),b.astimezone(timezone.utc)

def sales_query(a,b):
    q=select(Sale).where(Sale.created_at>=a,Sale.created_at<b)
    if g.user.role!='admin': q=q.where(Sale.user_id==g.user.id)
    return q

@bp.get('/')
@login_required
def dashboard():
    start,end,a,b=period()
    sales=g.db.scalars(sales_query(a,b).order_by(Sale.created_at)).all()
    valid=[s for s in sales if s.status=='completed']
    total=sum((s.total for s in valid),Decimal(0))
    days={str(start+timedelta(days=i)):Decimal(0) for i in range((end-start).days+1)}
    from ..security import aware
    for s in valid: days[str(aware(s.created_at).astimezone(TZ).date())]+=s.total
    low=g.db.scalars(select(Product).where(Product.active.is_(True),Product.stock<=Product.minimum).order_by(Product.stock).limit(10)).all() if g.user.role=='admin' else []
    return render_template('dashboard.html',start=start,end=end,total=total,count=len(valid),average=total/len(valid) if valid else 0,
        chart=[{'label':k[8:10]+'/'+k[5:7],'value':str(v)} for k,v in days.items()],low=low,recent=list(reversed(sales))[:10])

@bp.get('/sales')
@login_required
def sales():
    start,end,a,b=period(); q=sales_query(a,b)
    status=request.args.get('status','')
    if status in ('completed','cancelled'): q=q.where(Sale.status==status)
    rows=g.db.scalars(q.order_by(Sale.id.desc())).all()
    total=sum((s.total for s in rows if s.status=='completed'),Decimal(0))
    return render_template('sales.html',sales=rows,total=total,start=start,end=end,status=status)

@bp.get('/sales.xlsx')
@login_required
def export():
    start,end,a,b=period(); q=sales_query(a,b)
    status=request.args.get('status','')
    if status in ('completed','cancelled'): q=q.where(Sale.status==status)
    sales=g.db.scalars(q.order_by(Sale.id)).all()
    wb=Workbook(); ws=wb.active; ws.title='Vendas simuladas'
    ws.append(['Venda','Data São Paulo','Operador','Status','Total R$','Dinheiro','Débito','Crédito','Pix'])
    from ..security import aware
    for s in sales:
        methods={p.method:p.amount for p in s.payments}
        ws.append([s.id,aware(s.created_at).astimezone(TZ).strftime('%d/%m/%Y %H:%M'),s.user.name,s.status,float(s.total),*[float(methods.get(k,0)) for k in ('cash','debit','credit','pix')]])
        # Campos textuais são texto literal, inclusive nomes iniciados com =, +, - ou @.
        for cell in (ws.cell(ws.max_row,2),ws.cell(ws.max_row,3),ws.cell(ws.max_row,4)): cell.data_type='s'
    for cell in ws[1]: cell.font=Font(color='FFFFFF',bold=True); cell.fill=PatternFill('solid',fgColor='5630A6')
    for row in ws.iter_rows(min_row=2,min_col=5):
        for cell in row: cell.number_format='#,##0.00'
    for col,width in {'A':12,'B':22,'C':28,'D':18,'E':16,'F':16,'G':16,'H':16,'I':16}.items(): ws.column_dimensions[col].width=width
    ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
    out=BytesIO(); wb.save(out); out.seek(0)
    return send_file(out,as_attachment=True,download_name=f'vendas_{start}_{end}.xlsx',mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@bp.route('/finance',methods=['GET','POST'])
@admin_required
def finance():
    if request.method=='POST':
        reauthenticate()
        amount=number(request.form.get('amount',''),2,True)
        kind=request.form.get('kind')
        if kind not in ('income','expense'): raise ValueError('Tipo inválido.')
        desc=request.form.get('description','').strip()
        if not 5<=len(desc)<=180: raise ValueError('Descrição deve ter 5 a 180 caracteres.')
        g.db.add(Ledger(user_id=g.user.id,description='Manual: '+desc,amount=amount if kind=='income' else -amount))
        audit(g.db,g.user,'Lançamento financeiro manual'); g.db.commit()
        from flask import redirect,url_for
        return redirect(url_for('reports.finance'))
    start,end,a,b=period()
    rows=g.db.scalars(select(Ledger).where(Ledger.created_at>=a,Ledger.created_at<b).order_by(Ledger.id.desc())).all()
    opening=g.db.scalar(select(func.coalesce(func.sum(Ledger.amount),0)).where(Ledger.created_at<a))
    income=sum((r.amount for r in rows if r.amount>0),Decimal(0)); expense=-sum((r.amount for r in rows if r.amount<0),Decimal(0))
    return render_template('finance.html',rows=rows,start=start,end=end,opening=opening,income=income,expense=expense,balance=opening+income-expense)

@bp.get('/audit')
@admin_required
def audits():
    return render_template('audit.html',rows=g.db.scalars(select(Audit).order_by(Audit.id.desc()).limit(300)).all())
